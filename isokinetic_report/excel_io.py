from __future__ import annotations

import math
from pathlib import Path
from typing import Any, Iterable, Optional

import openpyxl
from openpyxl.worksheet.worksheet import Worksheet

from .models import (
    AsymmetryLevel,
    AthleteInfo,
    GaugeConfig,
    PriorityRule,
    Standards,
    TestRecord,
)


SHEET_ATHLETE = "1_运动员信息"
SHEET_DATA = "2_等速测试数据"
SHEET_STANDARDS = "3_绘图配置"
SHEET_COMMENTS = "4_原报告结论"

DATA_HEADERS = [
    "日期标签",
    "是否本次",
    "关节",
    "速度",
    "肌群A（比例分子）",
    "肌群B（比例分母）",
    "左A峰力矩_Nm",
    "左B峰力矩_Nm",
    "右A峰力矩_Nm",
    "右B峰力矩_Nm",
    "左比例_A/B",
    "右比例_A/B",
    "A双侧差异",
    "B双侧差异",
]

GAUGE_HEADERS = [
    "显示顺序",
    "关节",
    "仪表最小值",
    "低侧红区上界",
    "低侧橙区上界",
    "目标下限",
    "目标上限",
    "高侧黄区上界",
    "高侧橙区上界",
    "仪表最大值",
    "标准来源/状态",
    "启用",
    "备注",
]

ASYMMETRY_HEADERS = [
    "判定顺序",
    "状态",
    "最小值_含",
    "最大值_不含",
    "符号",
    "颜色HEX",
    "启用",
    "说明",
]

PRIORITY_HEADERS = [
    "规则组",
    "规则顺序",
    "输出标签",
    "指标",
    "比较符",
    "阈值",
    "组内关系",
    "标签颜色HEX",
    "启用",
    "说明",
]

PALETTE_HEADERS = ["用途", "颜色HEX", "说明"]

REQUIRED_PALETTE_KEYS = [
    "主色",
    "边框色",
    "目标范围",
    "轻度偏离",
    "中度偏离",
    "明显偏离",
    "缺失",
    "重点",
    "关注",
]

ALLOWED_PRIORITY_METRICS = {
    "比值红色项数",
    "比值偏离项数",
    "双侧差异↑项数",
    "双侧差异！或↑项数",
    "快速异常项数",
    "同速双侧比值异常次数",
    "总异常项数",
}


class WorkbookStructureError(ValueError):
    pass


def _text(value: Any) -> str:
    return "" if value is None else str(value).strip()


def _number(value: Any) -> Optional[float]:
    if value is None or value == "":
        return None
    try:
        number = float(value)
    except (TypeError, ValueError):
        return None
    return number if math.isfinite(number) else None


def _enabled(value: Any) -> bool:
    return _text(value).lower() in {"是", "启用", "true", "1", "yes"}


def _require_sheets(workbook: openpyxl.Workbook) -> None:
    missing = [name for name in [SHEET_ATHLETE, SHEET_DATA, SHEET_STANDARDS, SHEET_COMMENTS] if name not in workbook.sheetnames]
    if missing:
        raise WorkbookStructureError(f"缺少工作表：{', '.join(missing)}")


def _find_header_row(ws: Worksheet, required: Iterable[str], max_scan_rows: int = 80) -> Optional[int]:
    required_set = set(required)
    for row in range(1, min(ws.max_row, max_scan_rows) + 1):
        values = {_text(ws.cell(row, col).value) for col in range(1, ws.max_column + 1)}
        if required_set.issubset(values):
            return row
    return None


def _header_map(ws: Worksheet, row: int) -> dict[str, int]:
    return {_text(ws.cell(row, col).value): col for col in range(1, ws.max_column + 1) if ws.cell(row, col).value is not None}


def load_athlete_info(ws: Worksheet) -> AthleteInfo:
    header_row = _find_header_row(ws, ["字段", "值"])
    if header_row is None:
        raise WorkbookStructureError(f"{SHEET_ATHLETE} 中找不到“字段/值”表头")
    columns = _header_map(ws, header_row)
    values: dict[str, Any] = {}
    for row in range(header_row + 1, ws.max_row + 1):
        key = _text(ws.cell(row, columns["字段"]).value)
        if key:
            values[key] = ws.cell(row, columns["值"]).value
    required = ["姓名", "项目", "本次测试日期", "报告类型"]
    missing = [key for key in required if not _text(values.get(key))]
    if missing:
        raise WorkbookStructureError(f"{SHEET_ATHLETE} 缺少必填信息：{', '.join(missing)}")
    return AthleteInfo(
        name=_text(values["姓名"]),
        sport=_text(values["项目"]),
        sex=_text(values.get("性别")) or None,
        weight_kg=_number(values.get("体重_kg")),
        test_date=_text(values["本次测试日期"]),
        injury=_text(values.get("损伤情况")) or None,
        report_type=_text(values["报告类型"]),
    )


def load_test_data(ws_formula: Worksheet, ws_cached: Worksheet) -> list[TestRecord]:
    header_row = _find_header_row(ws_formula, DATA_HEADERS)
    if header_row is None:
        raise WorkbookStructureError(f"{SHEET_DATA} 缺少预期的 14 列表头")
    columns = _header_map(ws_formula, header_row)
    records: list[TestRecord] = []
    for row in range(header_row + 1, ws_formula.max_row + 1):
        if not any(ws_formula.cell(row, columns[h]).value is not None for h in DATA_HEADERS[:10]):
            continue
        get = lambda header: ws_formula.cell(row, columns[header]).value
        get_cached = lambda header: ws_cached.cell(row, columns[header]).value
        records.append(
            TestRecord(
                row_number=row,
                date_label=_text(get("日期标签")),
                is_current=_text(get("是否本次")) == "是",
                joint=_text(get("关节")),
                speed=_text(get("速度")),
                muscle_a=_text(get("肌群A（比例分子）")),
                muscle_b=_text(get("肌群B（比例分母）")),
                left_a=_number(get("左A峰力矩_Nm")),
                left_b=_number(get("左B峰力矩_Nm")),
                right_a=_number(get("右A峰力矩_Nm")),
                right_b=_number(get("右B峰力矩_Nm")),
                existing_left_ratio=_number(get_cached("左比例_A/B")),
                existing_right_ratio=_number(get_cached("右比例_A/B")),
                existing_a_asymmetry=_number(get_cached("A双侧差异")),
                existing_b_asymmetry=_number(get_cached("B双侧差异")),
            )
        )
    return records


def _missing_headers(ws: Worksheet, anchor_headers: list[str], full_headers: list[str]) -> tuple[Optional[int], list[str]]:
    row = _find_header_row(ws, anchor_headers)
    if row is None:
        return None, full_headers
    present = _header_map(ws, row)
    return row, [header for header in full_headers if header not in present]


def load_standards(ws: Worksheet) -> Standards:
    standards = Standards()

    gauge_row, gauge_missing = _missing_headers(ws, ["关节", "目标下限", "目标上限"], GAUGE_HEADERS)
    if gauge_row is not None:
        partial_columns = _header_map(ws, gauge_row)
        for row in range(gauge_row + 1, ws.max_row + 1):
            joint = _text(ws.cell(row, partial_columns["关节"]).value)
            if not joint:
                if standards.target_ranges:
                    break
                continue
            target_low = _number(ws.cell(row, partial_columns["目标下限"]).value)
            target_high = _number(ws.cell(row, partial_columns["目标上限"]).value)
            order = int(_number(ws.cell(row, partial_columns.get("显示顺序", 1)).value) or len(standards.target_ranges) + 1)
            if target_low is not None and target_high is not None and target_low < target_high:
                standards.target_ranges[joint] = (order, target_low, target_high)
    if gauge_missing:
        standards.issues.append(f"比例仪表盘配置缺少列：{', '.join(gauge_missing)}")
    elif gauge_row is not None:
        columns = _header_map(ws, gauge_row)
        for row in range(gauge_row + 1, ws.max_row + 1):
            joint = _text(ws.cell(row, columns["关节"]).value)
            if not joint:
                if standards.gauges:
                    break
                continue
            numeric_headers = GAUGE_HEADERS[2:10]
            numbers = {h: _number(ws.cell(row, columns[h]).value) for h in numeric_headers}
            missing_values = [h for h, value in numbers.items() if value is None]
            if missing_values:
                standards.issues.append(f"比例仪表盘配置第 {row} 行缺少数值：{', '.join(missing_values)}")
                continue
            sequence = [numbers[h] for h in numeric_headers]
            if not all(a < b for a, b in zip(sequence, sequence[1:])):
                standards.issues.append(f"比例仪表盘配置第 {row} 行边界必须严格递增")
                continue
            config = GaugeConfig(
                display_order=int(_number(ws.cell(row, columns["显示顺序"]).value) or 0),
                joint=joint,
                gauge_min=numbers["仪表最小值"],
                low_red_upper=numbers["低侧红区上界"],
                low_orange_upper=numbers["低侧橙区上界"],
                target_low=numbers["目标下限"],
                target_high=numbers["目标上限"],
                high_yellow_upper=numbers["高侧黄区上界"],
                high_orange_upper=numbers["高侧橙区上界"],
                gauge_max=numbers["仪表最大值"],
                source_status=_text(ws.cell(row, columns["标准来源/状态"]).value),
                enabled=_enabled(ws.cell(row, columns["启用"]).value),
                note=_text(ws.cell(row, columns["备注"]).value),
            )
            if config.enabled:
                standards.gauges[joint] = config
                standards.target_ranges[joint] = (config.display_order, config.target_low, config.target_high)

    asym_row, asym_missing = _missing_headers(ws, ["状态", "符号"], ASYMMETRY_HEADERS)
    if asym_missing:
        standards.issues.append(f"双侧差异配置缺少列：{', '.join(asym_missing)}")
    elif asym_row is not None:
        columns = _header_map(ws, asym_row)
        for row in range(asym_row + 1, ws.max_row + 1):
            state = _text(ws.cell(row, columns["状态"]).value)
            if not state:
                if standards.asymmetry_levels:
                    break
                continue
            if not _enabled(ws.cell(row, columns["启用"]).value):
                continue
            minimum = _number(ws.cell(row, columns["最小值_含"]).value)
            maximum = _number(ws.cell(row, columns["最大值_不含"]).value)
            if minimum is None:
                standards.issues.append(f"双侧差异配置第 {row} 行缺少最小值")
                continue
            standards.asymmetry_levels.append(
                AsymmetryLevel(
                    order=int(_number(ws.cell(row, columns["判定顺序"]).value) or 0),
                    state=state,
                    minimum_inclusive=minimum,
                    maximum_exclusive=maximum,
                    symbol=_text(ws.cell(row, columns["符号"]).value),
                    color=_text(ws.cell(row, columns["颜色HEX"]).value),
                    enabled=True,
                    note=_text(ws.cell(row, columns["说明"]).value),
                )
            )
        standards.asymmetry_levels.sort(key=lambda level: level.order)

    priority_row, priority_missing = _missing_headers(ws, ["规则组", "输出标签", "指标"], PRIORITY_HEADERS)
    if priority_missing:
        standards.issues.append(f"重点/关注规则缺少列：{', '.join(priority_missing)}")
    elif priority_row is not None:
        columns = _header_map(ws, priority_row)
        for row in range(priority_row + 1, ws.max_row + 1):
            group = _text(ws.cell(row, columns["规则组"]).value)
            if not group:
                if standards.priority_rules:
                    break
                continue
            if not _enabled(ws.cell(row, columns["启用"]).value):
                continue
            metric = _text(ws.cell(row, columns["指标"]).value)
            threshold = _number(ws.cell(row, columns["阈值"]).value)
            if metric not in ALLOWED_PRIORITY_METRICS:
                standards.issues.append(f"重点/关注规则第 {row} 行指标不可识别：{metric}")
                continue
            if threshold is None:
                standards.issues.append(f"重点/关注规则第 {row} 行缺少阈值")
                continue
            standards.priority_rules.append(
                PriorityRule(
                    group=group,
                    order=int(_number(ws.cell(row, columns["规则顺序"]).value) or 0),
                    output_label=_text(ws.cell(row, columns["输出标签"]).value),
                    metric=metric,
                    operator=_text(ws.cell(row, columns["比较符"]).value),
                    threshold=threshold,
                    group_relation=_text(ws.cell(row, columns["组内关系"]).value) or "任一",
                    label_color=_text(ws.cell(row, columns["标签颜色HEX"]).value),
                    enabled=True,
                    note=_text(ws.cell(row, columns["说明"]).value),
                )
            )

    palette_row, palette_missing = _missing_headers(ws, ["用途", "颜色HEX"], PALETTE_HEADERS)
    if palette_missing:
        standards.issues.append(f"颜色配置缺少列：{', '.join(palette_missing)}")
    elif palette_row is not None:
        columns = _header_map(ws, palette_row)
        for row in range(palette_row + 1, ws.max_row + 1):
            purpose = _text(ws.cell(row, columns["用途"]).value)
            if not purpose:
                if standards.palette:
                    break
                continue
            standards.palette[purpose] = _text(ws.cell(row, columns["颜色HEX"]).value)
        missing_keys = [key for key in REQUIRED_PALETTE_KEYS if not standards.palette.get(key)]
        if missing_keys:
            standards.issues.append(f"颜色配置缺少用途：{', '.join(missing_keys)}")

    if not standards.gauges and not gauge_missing:
        standards.issues.append("比例仪表盘配置没有启用的关节")
    if not standards.asymmetry_levels and not asym_missing:
        standards.issues.append("双侧差异配置没有启用的分级")
    if not standards.priority_rules and not priority_missing:
        standards.issues.append("重点/关注规则没有启用的规则")
    return standards


def load_original_comments(ws: Worksheet) -> dict[str, str]:
    header_row = _find_header_row(ws, ["部位", "原报告结论"])
    if header_row is None:
        return {}
    columns = _header_map(ws, header_row)
    comments: dict[str, str] = {}
    for row in range(header_row + 1, ws.max_row + 1):
        joint = _text(ws.cell(row, columns["部位"]).value)
        comment = _text(ws.cell(row, columns["原报告结论"]).value)
        if joint and comment:
            comments[joint] = comment
    return comments


def load_workbook_data(path: Path) -> tuple[AthleteInfo, list[TestRecord], Standards, dict[str, str]]:
    if not path.exists():
        raise FileNotFoundError(f"输入文件不存在：{path}")
    formula_book = openpyxl.load_workbook(path, data_only=False, read_only=False)
    cached_book = openpyxl.load_workbook(path, data_only=True, read_only=False)
    _require_sheets(formula_book)
    athlete = load_athlete_info(cached_book[SHEET_ATHLETE])
    records = load_test_data(formula_book[SHEET_DATA], cached_book[SHEET_DATA])
    standards = load_standards(cached_book[SHEET_STANDARDS])
    comments = load_original_comments(cached_book[SHEET_COMMENTS])
    formula_book.close()
    cached_book.close()
    return athlete, records, standards, comments
