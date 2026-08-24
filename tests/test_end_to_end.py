import re
import tempfile
import unittest
from datetime import datetime
from pathlib import Path

import openpyxl
from PIL import Image

from isokinetic_report.analysis import analyze
from isokinetic_report.excel_io import load_workbook_data
from isokinetic_report.render import generate_report


class EndToEndTests(unittest.TestCase):
    def _make_workbook(self, path: Path, joint_count: int = 5) -> None:
        wb = openpyxl.Workbook()
        ws1 = wb.active
        ws1.title = "1_运动员信息"
        ws1.append(["等速肌力报告输入文件"])
        ws1.append([])
        ws1.append(["字段", "值", "是否需手动输入", "说明"])
        for row in [
            ["姓名", "测试员", "是", ""], ["项目", "棒球", "是", ""], ["性别", "男", "是", ""],
            ["体重_kg", 80, "是", ""], ["本次测试日期", "2026-07-03", "是", ""], ["报告类型", "等速肌力测试", "否", ""],
        ]:
            ws1.append(row)

        ws2 = wb.create_sheet("2_等速测试数据")
        ws2.append(["等速测试原始数据"])
        ws2.append([])
        ws2.append(["日期标签", "是否本次", "关节", "速度", "肌群A（比例分子）", "肌群B（比例分母）", "左A峰力矩_Nm", "左B峰力矩_Nm", "右A峰力矩_Nm", "右B峰力矩_Nm", "左比例_A/B", "右比例_A/B", "A双侧差异", "B双侧差异"])
        joints = [
            ("肩关节屈伸", "屈肌", "伸肌"), ("肩关节内外旋", "外旋", "内旋"),
            ("髋关节屈伸", "屈肌", "伸肌"), ("膝关节屈伸", "屈肌", "伸肌"),
            ("踝关节屈伸", "背屈", "跖屈"), ("肘关节屈伸", "屈肌", "伸肌"),
            ("腕关节屈伸", "屈肌", "伸肌"), ("躯干屈伸", "屈肌", "伸肌"),
            ("颈部屈伸", "屈肌", "伸肌"), ("前臂旋转", "旋前", "旋后"),
        ][:joint_count]
        for index, (joint, a, b) in enumerate(joints):
            for speed in ["慢速", "快速"]:
                la, lb, ra, rb = 60 + index * 5, 100, 70 + index * 5, 105
                ws2.append(["7.3", "是", joint, speed, a, b, la, lb, ra, rb, la / lb, ra / rb, abs(la - ra) / max(la, ra), abs(lb - rb) / max(lb, rb)])

        ws3 = wb.create_sheet("3_绘图配置")
        ws3.append(["绘图配置"])
        ws3.append([])
        ws3.append(["显示顺序", "关节", "仪表最小值", "低侧红区上界", "低侧橙区上界", "目标下限", "目标上限", "高侧黄区上界", "高侧橙区上界", "仪表最大值", "标准来源/状态", "启用", "备注"])
        for order, (joint, _a, _b) in enumerate(joints, 1):
            ws3.append([order, joint, 0, 0.3, 0.5, 0.55, 0.75, 0.85, 1.0, 1.3, "测试标准", "是", ""])
        ws3.append([])
        ws3.append(["判定顺序", "状态", "最小值_含", "最大值_不含", "符号", "颜色HEX", "启用", "说明"])
        ws3.append([1, "正常", 0, 0.100000001, "√", "#159A36", "是", "双侧差异 ≤ 0.10"])
        ws3.append([2, "关注", 0.100000001, 0.2, "！", "#FF8A00", "是", "0.10 < 双侧差异 < 0.20"])
        ws3.append([3, "明显偏大", 0.2, None, "↑", "#F01818", "是", "双侧差异 ≥ 0.20"])
        ws3.append([])
        ws3.append(["规则组", "规则顺序", "输出标签", "指标", "比较符", "阈值", "组内关系", "标签颜色HEX", "启用", "说明"])
        ws3.append(["重点_慢速和快速异常", 1, "关注比值", "慢速和快速异常项数", ">=", 2, "任一", "#EF4444", "是", "存在多个明显比值异常"])
        ws3.append(["重点_明显差异", 2, "关注差异", "双侧差异↑项数", ">=", 2, "任一", "#EF4444", "是", "存在多个明显双侧差异"])
        ws3.append([])
        ws3.append(["用途", "颜色HEX", "说明"])
        for purpose, color in {
            "主色": "#123B8F", "边框色": "#A8C8FF", "目标范围": "#159A36", "轻度偏离": "#F5C518",
            "中度偏离": "#FF8A00", "明显偏离": "#F01818", "缺失": "#9AA4B2", "重点": "#F01818", "关注": "#FF8A00",
        }.items():
            ws3.append([purpose, color, ""])

        ws4 = wb.create_sheet("4_原报告结论")
        ws4.append(["原报告结论"])
        ws4.append([])
        ws4.append(["序号", "部位", "原报告结论"])
        for index, (joint, _a, _b) in enumerate(joints, 1):
            ws4.append([index, joint, "仅用于核对"])
        wb.save(path)

    def test_render_png_and_pdf(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            workbook_path = root / "input.xlsx"
            self._make_workbook(workbook_path)
            athlete, records, standards, comments = load_workbook_data(workbook_path)
            self.assertFalse(standards.issues)
            result = analyze(athlete, records, standards, comments)
            paths = generate_report(result, root / "output")
            self.assertTrue(paths.png.exists())
            self.assertTrue(paths.pdf.exists())
            with Image.open(paths.png) as image:
                self.assertEqual(image.size, (1600, 2000))
            self.assertEqual(paths.pdf.read_bytes()[:4], b"%PDF")

    def test_excel_datetime_is_normalized_to_calendar_date(self):
        with tempfile.TemporaryDirectory() as temp:
            workbook_path = Path(temp) / "datetime_input.xlsx"
            self._make_workbook(workbook_path)
            workbook = openpyxl.load_workbook(workbook_path)
            workbook["1_运动员信息"]["B8"] = datetime(2026, 8, 19, 14, 35, 20)
            workbook.save(workbook_path)
            athlete, _records, _standards, _comments = load_workbook_data(workbook_path)
            self.assertEqual(athlete.test_date, "2026-08-19")

    def test_three_joint_report_crops_unused_lower_space(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            workbook_path = root / "input_3.xlsx"
            self._make_workbook(workbook_path, joint_count=3)
            athlete, records, standards, comments = load_workbook_data(workbook_path)
            paths = generate_report(analyze(athlete, records, standards, comments), root / "output")
            with Image.open(paths.png) as image:
                self.assertEqual(image.size, (1600, 1560))

    def test_render_six_to_ten_joints_as_two_pages(self):
        for joint_count in (6, 10):
            with self.subTest(joint_count=joint_count), tempfile.TemporaryDirectory() as temp:
                root = Path(temp)
                workbook_path = root / f"input_{joint_count}.xlsx"
                self._make_workbook(workbook_path, joint_count=joint_count)
                athlete, records, standards, comments = load_workbook_data(workbook_path)
                result = analyze(athlete, records, standards, comments)
                paths = generate_report(result, root / "output")
                self.assertEqual(len(paths.pngs), 2)
                expected_sizes = [(1600, 2000), (1600, 1120)] if joint_count == 6 else [(1600, 2000), (1600, 2000)]
                for png, expected_size in zip(paths.pngs, expected_sizes):
                    with Image.open(png) as image:
                        self.assertEqual(image.size, expected_size)
                self.assertTrue(paths.pdf.exists())
                pdf_bytes = paths.pdf.read_bytes()
                self.assertEqual(len(re.findall(rb"/Type\s*/Page\b", pdf_bytes)), 2)


if __name__ == "__main__":
    unittest.main()
