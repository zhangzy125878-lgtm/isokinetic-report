import tempfile
import unittest
from pathlib import Path

import openpyxl
from PIL import Image

from isokinetic_report.analysis import analyze
from isokinetic_report.excel_io import load_workbook_data
from isokinetic_report.models import TestRecord
from isokinetic_report.render_peak_torque import _page_result, _report_joints, _torque_lines, generate_report
from run_report_peak_torque import _directory_excels, run as run_peak_torque
from tests import test_end_to_end


class PeakTorqueReportTests(unittest.TestCase):
    def test_directory_excels_returns_all_workbooks_and_skips_temporary_file(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            expected = [root / f"运动员{index}.xlsx" for index in range(1, 5)]
            for path in reversed(expected):
                path.touch()
            (root / "~$正在编辑.xlsx").touch()
            (root / "说明.txt").touch()
            self.assertEqual(_directory_excels(root), expected)

    def test_torque_lines_include_both_muscles_and_units(self):
        record = TestRecord(4, "7.3", True, "膝关节屈伸", "慢速", "屈肌", "伸肌", 99, 163, 69, 103)
        self.assertEqual(_torque_lines(record, "左侧"), ("屈 99 Nm", "伸 163 Nm"))
        self.assertEqual(_torque_lines(record, "右侧"), ("屈 69 Nm", "伸 103 Nm"))

    def test_single_muscle_torque_omits_empty_second_line(self):
        record = TestRecord(18, "8.10", True, "躯干旋转", "慢速", "旋转", "", 111, None, 105, None)
        self.assertEqual(_torque_lines(record, "左侧"), ("旋转 111 Nm", ""))

    def test_sheet_two_only_joint_is_appended_to_peak_torque_report(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            workbook_path = root / "input_4.xlsx"
            test_end_to_end.EndToEndTests()._make_workbook(workbook_path, joint_count=4)
            athlete, records, standards, comments = load_workbook_data(workbook_path)
            records.append(TestRecord(18, "8.10", True, "躯干旋转", "慢速", "旋转", "", 111, None, 105, None))
            result = analyze(athlete, records, standards, comments)
            self.assertEqual(_report_joints(result)[-1], ("躯干旋转", 5))

    def test_report_only_includes_joints_present_in_current_records(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            workbook_path = root / "input_4.xlsx"
            test_end_to_end.EndToEndTests()._make_workbook(workbook_path, joint_count=4)
            athlete, records, standards, comments = load_workbook_data(workbook_path)
            result = analyze(athlete, records, standards, comments)
            result.records = [
                record
                for record in result.records
                if record.joint in {"髋关节屈伸", "膝关节屈伸"}
            ]
            self.assertEqual(_report_joints(result), [("髋关节屈伸", 1), ("膝关节屈伸", 2)])

    def test_weaknesses_are_limited_to_the_current_page_joints(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            workbook_path = root / "input_6.xlsx"
            test_end_to_end.EndToEndTests()._make_workbook(workbook_path, joint_count=6)
            athlete, records, standards, comments = load_workbook_data(workbook_path)
            result = analyze(athlete, records, standards, comments)
            first_page = _page_result(result, {"关节1", "关节2", "关节3", "关节4", "关节5"})
            second_page = _page_result(result, {"关节6"})
            self.assertTrue(all("关节6" not in item for item in first_page.recommendations))
            self.assertTrue(all("关节6" in item for item in second_page.recommendations))

    def test_four_joint_peak_torque_report_uses_separate_taller_output(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            workbook_path = root / "input_4.xlsx"
            test_end_to_end.EndToEndTests()._make_workbook(workbook_path, joint_count=4)
            athlete, records, standards, comments = load_workbook_data(workbook_path)
            output_dir = root / "output"
            output_dir.mkdir()
            stale_page = output_dir / "测试员_等速肌力综合报告_20260703_峰力矩版_第1页.png"
            stale_page.touch()
            paths = generate_report(analyze(athlete, records, standards, comments), output_dir)
            self.assertIn("峰力矩版", paths.png.name)
            self.assertFalse(stale_page.exists())
            with Image.open(paths.png) as image:
                self.assertEqual(image.size, (1600, 1872))

    def test_batch_directory_generates_four_matching_reports(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            input_dir = root / "input"
            output_dir = root / "output"
            input_dir.mkdir()
            for index in range(1, 5):
                workbook_path = input_dir / f"运动员{index}.xlsx"
                test_end_to_end.EndToEndTests()._make_workbook(workbook_path, joint_count=1)
                workbook = openpyxl.load_workbook(workbook_path)
                workbook["1_运动员信息"]["B4"] = f"运动员{index}"
                workbook.save(workbook_path)
            exit_code = run_peak_torque([str(input_dir), "--output-dir", str(output_dir), "--no-pdf"])
            self.assertEqual(exit_code, 0)
            self.assertEqual(len(list(output_dir.glob("*.png"))), 4)
            self.assertEqual(len(list(output_dir.glob("*.pdf"))), 0)


if __name__ == "__main__":
    unittest.main()
